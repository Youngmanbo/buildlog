import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


# =========================
# 여기만 수정해서 사용
# =========================
IN_PATH = r"E:\250826 행정복합도시L-3 산출서.xlsx"
OUT_PATH = r"E:\ex_count.xlsx"

TARGET_SHEET_INDEX = 1   # 2번째 시트 (0=1번째, 1=2번째)
# TARGET_SHEET_NAME = "수량산출서"  # 이름으로 고정하고 싶으면 이걸 쓰고 아래에서 인덱스 대신 선택
# =========================


EA_PATTERN = re.compile(r"\[\s*\d+\s*EA\s*\]", re.IGNORECASE)

def norm_text(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\u00A0", " ")
    return re.sub(r"\s+", "", s)

def clean_token(s):
    if s is None:
        return None
    s = str(s).strip()
    s = EA_PATTERN.sub("", s).strip()  # [1EA] 제거
    s = re.sub(r"\s+", " ", s)
    return s

def extract_gongjong_value(row_values) -> str | None:
    texts = [str(v) if v is not None else "" for v in row_values]
    norms = [norm_text(v) for v in row_values]

    for i, nt in enumerate(norms):
        if "공종명" in nt:
            raw = texts[i].strip()

            # 같은 셀에 "공 종 명 : ..." 형태
            if ":" in raw:
                after = raw.split(":", 1)[1].strip()
                if after:
                    return after

            # 다음 셀에 값이 있는 경우
            if i + 1 < len(texts):
                nxt = texts[i + 1].strip()
                if nxt:
                    return nxt
            return None
    return None

def get_token(gongjong: str, idx0: int) -> str | None:
    if not gongjong:
        return None
    parts = [p.strip() for p in str(gongjong).split("_")]
    if len(parts) <= idx0:
        return None
    return parts[idx0]

def find_header_map(row_values) -> dict | None:
    norms = [norm_text(v) for v in row_values]
    targets = {"코드": None, "품명": None, "규격": None, "단위": None, "수량": None}

    for idx, nt in enumerate(norms):
        for key in list(targets.keys()):
            if targets[key] is None and key in nt:
                targets[key] = idx

    if all(v is not None for v in targets.values()):
        return targets
    return None

def to_float(x) -> float:
    if x is None:
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace(",", "")
    if not s:
        return 0.0
    try:
        return float(s)
    except ValueError:
        return 0.0

def parse_one_sheet(input_path: str) -> pd.DataFrame:
    wb = load_workbook(input_path, data_only=True)

    ws_list = wb.worksheets
    if TARGET_SHEET_INDEX < 0 or TARGET_SHEET_INDEX >= len(ws_list):
        raise IndexError(f"시트 인덱스 오류: {TARGET_SHEET_INDEX}, 전체={len(ws_list)}")

    ws = ws_list[TARGET_SHEET_INDEX]
    # ws = wb[TARGET_SHEET_NAME]  # 이름으로 선택하고 싶으면 이 줄을 사용

    records = []
    current_gongjong = None
    key4 = None            # 4번째 토큰(세대라인)
    work5 = None           # 5번째 토큰(공사종류)
    header_map = None
    in_table = False

    max_col = ws.max_column

    for r in range(1, ws.max_row + 1):
        row_values = [ws.cell(r, c).value for c in range(1, max_col + 1)]

        # 공종명 갱신(페이지 시작)
        gj = extract_gongjong_value(row_values)
        if gj:
            current_gongjong = gj
            key4 = clean_token(get_token(current_gongjong, 3))   # 4번째
            work5 = clean_token(get_token(current_gongjong, 4))  # 5번째 = 공사종류
            header_map = None
            in_table = False

        # 헤더 찾기
        hm = find_header_map(row_values)
        if hm:
            header_map = hm
            in_table = True
            continue

        # 데이터 읽기
        if in_table and header_map and key4 and work5:
            code = row_values[header_map["코드"]]
            name = row_values[header_map["품명"]]
            spec = row_values[header_map["규격"]]
            unit = row_values[header_map["단위"]]
            qty  = row_values[header_map["수량"]]

            # 빈 행 스킵
            if code is None and name is None and spec is None and unit is None and qty is None:
                continue

            code_s = "" if code is None else str(code).strip().lstrip("'")
            qty_f = to_float(qty)

            # 필터링
            if qty_f <= 0:
                continue
            if "ZZZZ" in code_s:
                continue

            records.append({
                "세대라인": key4,
                "공사종류": work5,  # ✅ 추가된 컬럼
                "코드": code_s,
                "품명": "" if name is None else str(name).strip(),
                "규격": "" if spec is None else str(spec).strip(),
                "단위": "" if unit is None else str(unit).strip(),
                "수량": qty_f,
            })

    return pd.DataFrame(records)

def run():
    if not Path(IN_PATH).exists():
        raise FileNotFoundError(f"입력 파일이 없습니다: {IN_PATH}")

    df = parse_one_sheet(IN_PATH)

    # ✅ 공사종류까지 포함해서 집계 (서로 다른 공사종류가 섞여 합산되는 것 방지)
    result = (
        df.groupby(["세대라인", "공사종류", "코드", "품명", "규격", "단위"], as_index=False)["수량"]
          .sum()
          .sort_values(["세대라인", "공사종류", "코드"])
    )

    Path(OUT_PATH).parent.mkdir(parents=True, exist_ok=True)
    result.to_excel(OUT_PATH, index=False)

    print("완료:", OUT_PATH)
    print("원시 행:", len(df), " / 집계 행:", len(result))

if __name__ == "__main__":
    run()
